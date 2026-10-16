"""Genera los documentos editables y de presentación del Gate 0."""

from pathlib import Path
import subprocess
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches, Pt as PptPt
import markdown


ROOT = Path(__file__).resolve().parent
NAVY = "10283F"
TEAL = "087E8B"
CYAN = "42B7C8"
MINT = "D8F0EA"
GOLD = "E9AD45"
CORAL = "C9584F"
INK = "263746"
MUTED = "627482"
PAPER = "F4F7F8"
WHITE = "FFFFFF"
LINE = "DCE5E9"


def rgb(value: str) -> PptColor:
    return PptColor.from_string(value)


def shape(slide, kind, x, y, w, h, fill, line=None):
    obj = slide.shapes.add_shape(kind, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    obj.line.fill.background() if line is None else None
    if line:
        obj.line.color.rgb = rgb(line)
    return obj


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font="Aptos",
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = PptInches(margin)
    tf.margin_top = tf.margin_bottom = PptInches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_footer(slide, number, dark=False):
    color = "B7C8D3" if dark else MUTED
    shape(slide, MSO_SHAPE.RECTANGLE, 0.55, 7.08, 12.2, 0.012, "315168" if dark else LINE)
    text(slide, "PREFactibilidad  •  16 OCT 2026", 0.6, 7.12, 5, 0.2, 8, color)
    text(slide, f"RÍO BLANCO  /  {number:02d}", 10.4, 7.12, 2.3, 0.2, 8, color, align=PP_ALIGN.RIGHT)


def new_slide(prs, title=None, kicker=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = NAVY if dark else PAPER
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    if title:
        text(slide, kicker.upper() if kicker else "MCH RÍO BLANCO  •  PUNO", 0.65, 0.38, 11.8, 0.25, 9, CYAN if dark else TEAL, True)
        text(slide, title, 0.65, 0.78, 12.0, 0.65, 27, WHITE if dark else NAVY, True)
    add_footer(slide, len(prs.slides), dark)
    return slide


def card(slide, x, y, w, h, title, body, accent=TEAL, dark=False, body_size=14):
    fill = "19354D" if dark else WHITE
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.055, h, accent)
    text(slide, title, x + 0.22, y + 0.18, w - 0.4, 0.4, 16, WHITE if dark else NAVY, True)
    text(slide, body, x + 0.22, y + 0.68, w - 0.42, h - 0.84, body_size, "D6E3EA" if dark else INK)


def add_presentation(path: Path):
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    # 1 — cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, 7.5, TEAL)
    # Abstract river / penstock motif, deliberately schematic.
    for i, (x, y, w, h, col) in enumerate([(9.4, 0.9, 2.8, 2.8, "193B54"), (10.2, 1.7, 1.8, 1.8, "20516A"), (10.75, 2.25, 0.7, 0.7, CYAN)]):
        shape(s, MSO_SHAPE.OVAL, x, y, w, h, col)
    shape(s, MSO_SHAPE.RECTANGLE, 9.95, 3.55, 0.07, 2.2, CYAN)
    shape(s, MSO_SHAPE.RECTANGLE, 9.95, 5.65, 2.05, 0.07, CYAN)
    text(s, "DECISIÓN DE INVERSIÓN", 0.8, 0.7, 6.8, 0.35, 12, CYAN, True)
    text(s, "Microcentral\nhidroeléctrica", 0.8, 1.35, 7.8, 1.65, 34, WHITE, True)
    text(s, "Río Blanco  /  100 kW  /  Ollachea, Puno", 0.82, 3.3, 7.4, 0.5, 18, "D4E3EA")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, 4.35, 6.6, 0.9, "19354D")
    text(s, "¿Conviene financiar estudios de campo?", 1.08, 4.55, 6.1, 0.45, 20, WHITE, True)
    text(s, "Prefactibilidad  •  Paquete para revisión ejecutiva  •  16 octubre 2026", 0.83, 6.55, 8.8, 0.3, 10, "B7C8D3")
    add_footer(s, 1, True)

    # 2 — decision
    s = new_slide(prs, "La decisión de hoy es comprar información", "01 / recomendación")
    card(s, 0.72, 1.75, 5.85, 2.0, "Recomendación", "Autorizar solo Gate 0, con tope de USD 25,000 y entregables de parada. No comprometer aún el presupuesto de factibilidad ni construcción.", TEAL, body_size=17)
    card(s, 6.78, 1.75, 5.8, 2.0, "Por qué ahora", "El caso financiero cambia de positivo a negativo bajo estrés combinado. Topografía, caudal turbinable y CAPEX deben verificarse antes de decidir la campaña completa.", GOLD, body_size=17)
    text(s, "39", 0.9, 4.45, 1.6, 1.1, 50, NAVY, True)
    text(s, "/100", 2.1, 4.88, 1.4, 0.45, 20, MUTED, True)
    text(s, "Calificación del expediente disponible", 3.4, 4.73, 4.5, 0.45, 17, NAVY, True)
    text(s, "La nota refleja falta de evidencia primaria; no mide el potencial real del río.", 3.4, 5.25, 7.9, 0.7, 13, MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, 6.2, 11.65, 0.48, MINT)
    text(s, "Gate 0 = una decisión escalonada, reversible y limitada", 1.05, 6.29, 11.2, 0.25, 14, NAVY, True)

    # 3 — baseline
    s = new_slide(prs, "El caso base luce atractivo, pero usa supuestos", "02 / caso reproducido")
    metrics = [("100 kW", "potencia nominal"), ("115 m", "salto bruto asumido"), ("125 L/s", "caudal de diseño"), ("USD 245.8k", "CAPEX supuesto")]
    for i, (v, lab) in enumerate(metrics):
        x = 0.75 + i * 3.1
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.75, 2.82, 1.25, WHITE)
        text(s, v, x + 0.15, 1.95, 2.52, 0.48, 25, TEAL, True, align=PP_ALIGN.CENTER)
        text(s, lab, x + 0.1, 2.52, 2.62, 0.25, 11, MUTED, align=PP_ALIGN.CENTER)
    result = [("862.9", "MWh/año"), ("274.7k", "VAN USD"), ("23.2%", "TIR"), ("1.88", "B/C")]
    for i, (v, lab) in enumerate(result):
        x = 0.75 + i * 3.1
        text(s, v, x + 0.12, 3.7, 2.58, 0.58, 28, NAVY, True, align=PP_ALIGN.CENTER)
        text(s, lab, x + 0.1, 4.34, 2.62, 0.28, 12, MUTED, align=PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, 5.25, 11.62, 0.95, "FFF0D1")
    text(s, "Estos indicadores son resultados del modelo reproducible, no una previsión ni una oferta bancaria.", 1.08, 5.52, 11.1, 0.38, 16, NAVY, True)
    text(s, "Serie 2014–2023 sin fuente primaria acreditada; eficiencia global, tarifa, CAPEX y caudal ecológico asumidos.", 0.92, 6.42, 11.6, 0.28, 10, MUTED)

    # 4 — stress chart
    s = new_slide(prs, "El estrés combinado vuelve negativo el VAN", "03 / prueba de estrés")
    vals = [274.7, 211.8, 80.4, 180.9, -64.7]
    labels = ["Base", "Hb 100 m", "Q diseño\n80 L/s", "CAPEX\n+30%", "Combinado"]
    chart_x, chart_y, chart_w, chart_h = 1.25, 1.9, 10.8, 3.7
    min_v, max_v = -100, 300
    zero_y = chart_y + chart_h * max_v / (max_v - min_v)
    shape(s, MSO_SHAPE.RECTANGLE, chart_x, zero_y, chart_w, 0.025, MUTED)
    for tick in [300, 200, 100, 0, -100]:
        yy = chart_y + chart_h * (max_v - tick) / (max_v - min_v)
        text(s, f"{tick:+d}" if tick < 0 else str(tick), 0.65, yy - 0.1, 0.45, 0.2, 9, MUTED, align=PP_ALIGN.RIGHT)
        if tick != 0:
            shape(s, MSO_SHAPE.RECTANGLE, chart_x, yy, chart_w, 0.008, LINE)
    gap = chart_w / len(vals)
    for i, (v, lab) in enumerate(zip(vals, labels)):
        cx = chart_x + i * gap + gap / 2
        top_value_y = chart_y + chart_h * (max_v - max(v, 0)) / (max_v - min_v)
        if v >= 0:
            bh = (v / (max_v - min_v)) * chart_h
            by = zero_y - bh
        else:
            bh = (abs(v) / (max_v - min_v)) * chart_h
            by = zero_y
        col = CORAL if v < 0 else (TEAL if i < 4 else GOLD)
        shape(s, MSO_SHAPE.RECTANGLE, cx - 0.47, by, 0.94, bh, col)
        label_y = by - 0.34 if v >= 0 else by + bh + 0.02
        text(s, f"{v:+.1f}" if v < 0 else f"{v:.1f}", cx - 0.65, label_y, 1.3, 0.25, 11, NAVY, True, align=PP_ALIGN.CENTER)
        text(s, lab, cx - 0.85, 5.82, 1.7, 0.5, 10, INK, align=PP_ALIGN.CENTER)
    text(s, "VAN (miles de USD)  •  25 años  •  descuento 10%  •  supuestos del escenario vigente", 1.28, 6.5, 10.7, 0.25, 10, MUTED)

    # 5 — flow interpretation
    s = new_slide(prs, "80 L/s en estiaje no significa 80 L/s todo el año", "04 / lectura hidrológica")
    card(s, 0.8, 1.75, 5.75, 2.15, "Proxy de caudal de diseño", "Se limita la turbina a 80 L/s manteniendo la serie hipotética. VAN combinado con Hb 100 m y CAPEX +30%: −USD 64.7k.", GOLD, body_size=16)
    card(s, 6.78, 1.75, 5.75, 2.15, "Límite severo de estrés", "Se supone 80 L/s en el río todo el año. Con 30 L/s reservados quedan 50 L/s turbinables; VAN combinado: −USD 192.9k.", CORAL, body_size=16)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.83, 4.45, 11.7, 1.15, "E5EFF3")
    text(s, "Para modelar el estiaje real faltan la duración del caudal, una curva de duración validada y el caudal ecológico aplicable.", 1.12, 4.73, 11.1, 0.5, 16, NAVY, True)
    text(s, "El escenario constante es un límite de estrés, no una predicción hidrológica.", 1.12, 5.9, 11.1, 0.4, 12, MUTED)

    # 6 — scope
    s = new_slide(prs, "Gate 0: señales tempranas; campañas ampliadas opcionales", "05 / procura")
    card(s, 0.75, 1.72, 3.85, 3.85, "A  •  Topografía", "Revisión y reconocimiento\nControl de campo\nHasta 8 ha para cotizar\nPolígono real antes de movilizar\n\nEntregables: puntos, perfiles, secciones, crudos y CRS", TEAL, body_size=13)
    card(s, 4.75, 1.72, 3.85, 3.85, "B  •  Hidrología", "Auditoría de series y fuentes\nReconocimiento de sección\n2–4 aforos iniciales\n\nOpción ampliada: 12 meses y 15–20 aforos; crecidas y disponibilidad con incertidumbre", GOLD, body_size=13)
    card(s, 8.75, 1.72, 3.85, 3.85, "C  •  Geotecnia", "Revisión y recorrido de sitio\nPrograma de exploración cotizado\n\nOpción ampliada: 14 calicatas y 3 perforaciones iniciales, ajustadas por el geotecnista", CORAL, body_size=13)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.1, 5.95, 11.1, 0.58, MINT)
    text(s, "Pedir precio por lote, etapa, unidad, IGV, movilización, plazo, entregables y vigencia.", 1.35, 6.08, 10.6, 0.3, 14, NAVY, True, align=PP_ALIGN.CENTER)

    # 7 — constraints
    s = new_slide(prs, "Las restricciones todavía no están despejadas", "06 / matriz de control")
    cards = [
        ("Hidráulica", "Hb < 95 m: revaluar diseño y VAN", TEAL),
        ("Agua", "Aforo < 70 L/s: alerta; validar duración", GOLD),
        ("Predial", "Sin corredor/servidumbre: explorar variante", CORAL),
        ("Geotecnia", "Peligro activo: estimar mitigación", TEAL),
        ("Conexión", "Denegación: costear alternativas", GOLD),
        ("Ambiental", "Clasificación y permisos pendientes", CORAL),
    ]
    for i, (title, body, col) in enumerate(cards):
        x = 0.8 + (i % 3) * 4.05
        y = 1.75 + (i // 3) * 1.75
        card(s, x, y, 3.72, 1.42, title, body, col, body_size=12)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.86, 5.65, 11.75, 0.7, "FFF0D1")
    text(s, "Los límites 95 m / 70 L/s son disparadores de revisión, no cancelación automática.", 1.08, 5.86, 11.3, 0.3, 14, NAVY, True)
    text(s, "Estado actual: pendientes por falta de polígono verificado, datos de campo y respuestas institucionales.", 0.95, 6.48, 11.5, 0.28, 10, MUTED)

    # 8 — staged budget
    s = new_slide(prs, "Financiar por etapas, con ofertas antes de avanzar", "07 / capital de desarrollo")
    stages = [("GATE 0", "USD 10–25k", "screening + reconocimiento", TEAL),
              ("GATE 1", "USD 30–60k", "factibilidad de sitio", GOLD),
              ("GATE 2", "USD 60–110k", "factibilidad cerrada / detalle", CORAL)]
    for i, (name, amount, scope, col) in enumerate(stages):
        x = 0.8 + i * 4.05
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, 3.65, 2.45, WHITE)
        shape(s, MSO_SHAPE.RECTANGLE, x, 2.0, 3.65, 0.12, col)
        text(s, name, x + 0.2, 2.35, 3.25, 0.32, 13, col, True, align=PP_ALIGN.CENTER)
        text(s, amount, x + 0.2, 2.95, 3.25, 0.62, 24, NAVY, True, align=PP_ALIGN.CENTER)
        text(s, scope, x + 0.25, 3.75, 3.15, 0.38, 12, MUTED, align=PP_ALIGN.CENTER)
        if i < 2:
            text(s, "→", x + 3.67, 2.95, 0.35, 0.4, 23, MUTED, True, align=PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 5.08, 11.7, 1.18, "E5EFF3")
    text(s, "Rangos ROM para planificación; no son ofertas ni deben sumarse sin revisar solapes.", 1.12, 5.35, 11.15, 0.32, 16, NAVY, True)
    text(s, "La campaña hidrológica anual puede exigir plazo y fondos adicionales a Gate 0.", 1.12, 5.82, 11.1, 0.28, 12, MUTED)

    # 9 — what approval buys
    s = new_slide(prs, "Aprobar Gate 0 con entregables y criterios de salida", "08 / solicitud ejecutiva", dark=True)
    text(s, "AUTORIZAR", 0.9, 1.72, 3.0, 0.3, 11, CYAN, True)
    text(s, "Hasta USD 25,000", 0.9, 2.12, 6.4, 0.7, 31, WHITE, True)
    text(s, "para obtener ofertas comparables, fijar polígono de trabajo, ejecutar reconocimiento y resolver incertidumbres de mayor impacto.", 0.92, 3.0, 5.8, 1.25, 17, "D6E3EA")
    card(s, 7.05, 1.8, 5.2, 1.25, "1  Cotizar", "Tres lotes; base y opciones por separado.", TEAL, dark=True, body_size=13)
    card(s, 7.05, 3.3, 5.2, 1.25, "2  Verificar", "Geometría, agua, acceso, predios y conexión preliminar.", GOLD, dark=True, body_size=13)
    card(s, 7.05, 4.8, 5.2, 1.25, "3  Decidir", "Pasar a Gate 1 solo con ruta plausible y costos actualizados.", CORAL, dark=True, body_size=13)

    # 10 — close
    s = new_slide(prs, "Siguiente paso: cotizar y confirmar la ruta", "09 / acción inmediata")
    text(s, "1", 0.9, 1.7, 0.6, 0.65, 28, TEAL, True)
    text(s, "Enviar el TDR a postores de topografía, hidrología y geotecnia.", 1.65, 1.78, 9.8, 0.42, 17, NAVY, True)
    text(s, "2", 0.9, 2.75, 0.6, 0.65, 28, TEAL, True)
    text(s, "Pedir cotización Gate 0 y opciones ampliadas por separado.", 1.65, 2.83, 9.8, 0.42, 17, NAVY, True)
    text(s, "3", 0.9, 3.8, 0.6, 0.65, 28, TEAL, True)
    text(s, "Aprobar eje/polígono, acceso y responsable de recibir datos crudos.", 1.65, 3.88, 9.8, 0.42, 17, NAVY, True)
    text(s, "4", 0.9, 4.85, 0.6, 0.65, 28, TEAL, True)
    text(s, "Revisar ofertas, completar matriz y ejecutar el primer go/no-go.", 1.65, 4.93, 9.8, 0.42, 17, NAVY, True)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 6.02, 11.75, 0.6, NAVY)
    text(s, "Todavía sin levantamiento real • sin aforo de campo • sin overlay GIS confirmado", 1.08, 6.17, 11.25, 0.25, 13, WHITE, True, align=PP_ALIGN.CENTER)

    prs.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Pág. ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_inline(paragraph, node, bold=False, italic=False, code=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        return
    if not isinstance(node, Tag):
        return
    b = bold or node.name in {"strong", "b"}
    i = italic or node.name in {"em", "i"}
    c = code or node.name == "code"
    if node.name == "a":
        for child in node.children:
            add_inline(paragraph, child, b, i, c)
        href = node.get("href")
        if href:
            r = paragraph.add_run(f" ({href})")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(TEAL)
        return
    for child in node.children:
        add_inline(paragraph, child, b, i, c)


def md_to_docx(md_path: Path, docx_path: Path, label: str):
    html = markdown.markdown(md_path.read_text(encoding="utf-8"), extensions=["tables", "fenced_code", "sane_lists"])
    soup = BeautifulSoup(html, "html.parser")
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in [("Title", 25, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, NAVY)]:
        st = styles[name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = f"MCH RÍO BLANCO  /  {label.upper()}"
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    header.runs[0].font.bold = True
    footer = sec.footer.paragraphs[0]
    footer.text = "PREFactibilidad • Documento de trabajo • 16 octubre 2026"
    footer.style = styles["Caption"]
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(sec.footer.add_paragraph())

    body = soup.body or soup
    for elem in body.children:
        if not isinstance(elem, Tag):
            continue
        if elem.name in {"h1", "h2", "h3", "h4"}:
            level = int(elem.name[1])
            p = doc.add_paragraph(style="Title" if level == 1 else f"Heading {min(level,3)}")
            add_inline(p, elem)
        elif elem.name == "p":
            p = doc.add_paragraph()
            if elem.find("strong") and ("Estado" in elem.get_text() or "Advertencia" in elem.get_text()):
                p.paragraph_format.left_indent = Inches(0.12)
                p.paragraph_format.right_indent = Inches(0.1)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EAF2F4"); pPr.append(shd)
            add_inline(p, elem)
        elif elem.name in {"ul", "ol"}:
            style = "List Bullet" if elem.name == "ul" else "List Number"
            for li in elem.find_all("li", recursive=False):
                p = doc.add_paragraph(style=style)
                add_inline(p, li)
        elif elem.name == "table":
            rows = elem.find_all("tr")
            if not rows:
                continue
            ncols = max(len(r.find_all(["th", "td"], recursive=False)) for r in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for ridx, tr in enumerate(rows):
                cells = table.add_row().cells
                for cidx, td in enumerate(tr.find_all(["th", "td"], recursive=False)):
                    cells[cidx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cells[cidx].paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    add_inline(p, td)
                    for run in p.runs:
                        run.font.size = Pt(7.4)
                        if ridx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string(WHITE)
                    if ridx == 0:
                        set_cell_shading(cells[cidx], NAVY)
                    elif ridx % 2 == 0:
                        set_cell_shading(cells[cidx], "F0F5F6")
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif elem.name == "blockquote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.2)
            add_inline(p, elem)
        elif elem.name == "pre":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            add_inline(p, elem)
            for run in p.runs:
                run.font.name = "Consolas"; run.font.size = Pt(8)
    doc.save(docx_path)


def add_pdf_presentation(path: Path):
    """Rendera una versión PDF breve del deck, con el mismo sistema visual."""
    page_w, page_h = 960, 540
    canvas = Canvas(str(path), pagesize=(page_w, page_h), pageCompression=1)
    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    bold_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    if font_path.exists() and bold_path.exists():
        pdfmetrics.registerFont(TTFont("MCHSans", str(font_path)))
        pdfmetrics.registerFont(TTFont("MCHSans-Bold", str(bold_path)))
        regular, bold = "MCHSans", "MCHSans-Bold"
    else:
        regular, bold = "Helvetica", "Helvetica-Bold"

    def col(hexcolor):
        return HexColor("#" + hexcolor)

    def wrapped(value, x, top, width, size=16, color=INK, is_bold=False, leading=None, max_lines=None):
        font = bold if is_bold else regular
        leading = leading or size * 1.3
        lines = []
        for para in value.split("\n"):
            words = para.split()
            line = ""
            for word in words:
                candidate = f"{line} {word}".strip()
                if pdfmetrics.stringWidth(candidate, font, size) <= width or not line:
                    line = candidate
                else:
                    lines.append(line); line = word
            lines.append(line)
        if max_lines and len(lines) > max_lines:
            lines = lines[:max_lines]
            lines[-1] = lines[-1].rstrip(" .") + "…"
        canvas.setFont(font, size); canvas.setFillColor(col(color))
        y = page_h - top - size
        for line in lines:
            canvas.drawString(x, y, line)
            y -= leading
        return y

    def box(x, top, w, h, fill, radius=12, accent=None):
        y = page_h - top - h
        canvas.setFillColor(col(fill)); canvas.roundRect(x, y, w, h, radius, fill=1, stroke=0)
        if accent:
            canvas.setFillColor(col(accent)); canvas.roundRect(x, y, 5, h, 2, fill=1, stroke=0)

    def card(x, top, w, h, title, body, accent=TEAL, dark=False, size=14):
        box(x, top, w, h, "19354D" if dark else WHITE, accent=accent)
        wrapped(title, x + 18, top + 16, w - 36, 16, WHITE if dark else NAVY, True)
        wrapped(body, x + 18, top + 48, w - 36, size, "D6E3EA" if dark else INK, leading=size*1.35)

    page = 0
    def footer(dark=False):
        nonlocal page
        page += 1
        c = "B7C8D3" if dark else MUTED
        canvas.setStrokeColor(col("315168" if dark else LINE)); canvas.setLineWidth(0.8)
        canvas.line(42, 30, page_w - 42, 30)
        canvas.setFont(regular, 8); canvas.setFillColor(col(c))
        canvas.drawString(46, 16, "PREFACTIBILIDAD  •  16 OCT 2026")
        canvas.drawRightString(page_w - 46, 16, f"RÍO BLANCO  /  {page:02d}")
        canvas.showPage()

    def start(title=None, kicker=None, dark=False):
        canvas.setFillColor(col(NAVY if dark else PAPER)); canvas.rect(0, 0, page_w, page_h, fill=1, stroke=0)
        if title:
            wrapped((kicker or "MCH RÍO BLANCO  •  PUNO").upper(), 48, 28, 850, 9, CYAN if dark else TEAL, True)
            wrapped(title, 48, 60, 860, 27, WHITE if dark else NAVY, True, leading=30, max_lines=1)

    # Slide 1: cover
    start(dark=True)
    canvas.setFillColor(col(TEAL)); canvas.rect(0, 0, 14, page_h, fill=1, stroke=0)
    for x, y, r, color in [(720, 400, 108, "193B54"), (770, 350, 68, "20516A"), (808, 312, 27, CYAN)]:
        canvas.setFillColor(col(color)); canvas.circle(x, y, r, fill=1, stroke=0)
    canvas.setFillColor(col(CYAN)); canvas.roundRect(736, 125, 5, 170, 2, fill=1, stroke=0); canvas.roundRect(736, 123, 150, 5, 2, fill=1, stroke=0)
    wrapped("DECISIÓN DE INVERSIÓN", 58, 62, 570, 12, CYAN, True)
    wrapped("Microcentral\nhidroeléctrica", 58, 112, 650, 35, WHITE, True, 42)
    wrapped("Río Blanco  /  100 kW  /  Ollachea, Puno", 60, 260, 620, 18, "D4E3EA")
    box(60, 330, 520, 66, "19354D", radius=12)
    wrapped("¿Conviene financiar estudios de campo?", 82, 350, 485, 19, WHITE, True)
    wrapped("Prefactibilidad • Paquete para revisión ejecutiva • 16 octubre 2026", 60, 465, 820, 10, "B7C8D3")
    footer(True)

    # Slide 2: recommendation
    start("La decisión de hoy es comprar información", "01 / recomendación")
    card(58, 130, 408, 145, "Recomendación", "Autorizar solo Gate 0, con tope de USD 25,000 y entregables de parada. No comprometer aún factibilidad ni construcción.", TEAL, size=14)
    card(494, 130, 408, 145, "Por qué ahora", "El caso cambia de VAN positivo a negativo bajo estrés combinado. Topografía, caudal turbinable y CAPEX deben verificarse.", GOLD, size=14)
    wrapped("39", 65, 310, 120, 50, NAVY, True)
    wrapped("/100", 177, 335, 110, 21, MUTED, True)
    wrapped("Calificación del expediente disponible", 295, 328, 500, 17, NAVY, True)
    wrapped("La nota refleja falta de evidencia primaria; no mide el potencial real del río.", 295, 365, 550, 12, MUTED)
    box(62, 425, 840, 43, MINT, radius=9)
    wrapped("Gate 0 = una decisión escalonada, reversible y limitada", 80, 437, 800, 14, NAVY, True)
    footer()

    # Slide 3: base scenario
    start("El caso base luce atractivo, pero usa supuestos", "02 / caso reproducido")
    metrics = [("100 kW", "potencia nominal"), ("115 m", "salto bruto asumido"), ("125 L/s", "caudal diseño"), ("USD 245.8k", "CAPEX supuesto")]
    for i, (v, lab) in enumerate(metrics):
        x = 56 + i * 218
        box(x, 135, 198, 88, WHITE)
        wrapped(v, x + 8, 153, 182, 24, TEAL, True)
        wrapped(lab, x + 8, 190, 182, 10, MUTED)
    results = [("862.9", "MWh/año"), ("274.7k", "VAN USD"), ("23.2%", "TIR"), ("1.88", "B/C")]
    for i, (v, lab) in enumerate(results):
        x = 56 + i * 218
        wrapped(v, x, 270, 198, 29, NAVY, True)
        wrapped(lab, x, 312, 198, 11, MUTED)
    box(62, 365, 840, 75, "FFF0D1", radius=10)
    wrapped("Indicadores reproducibles del modelo; no son previsión ni oferta bancaria.", 82, 387, 800, 15, NAVY, True)
    wrapped("Serie 2014–2023 sin fuente primaria acreditada; eficiencia, tarifa, CAPEX y caudal ecológico asumidos.", 68, 454, 830, 9, MUTED)
    footer()

    # Slide 4: stress chart
    start("El estrés combinado vuelve negativo el VAN", "03 / prueba de estrés")
    vals = [274.7, 211.8, 80.4, 180.9, -64.7]
    labels = ["Base", "Hb 100 m", "Q diseño\n80 L/s", "CAPEX\n+30%", "Combinado"]
    cx0, cy0, cw, ch = 105, 135, 740, 270
    min_v, max_v = -100, 300
    zero = page_h - (cy0 + ch * max_v / (max_v - min_v))
    canvas.setStrokeColor(col(MUTED)); canvas.line(cx0, zero, cx0 + cw, zero)
    for tick in [300, 200, 100, 0, -100]:
        yy = page_h - (cy0 + ch * (max_v - tick) / (max_v - min_v))
        canvas.setStrokeColor(col(LINE)); canvas.line(cx0, yy, cx0 + cw, yy)
        wrapped(str(tick), 60, page_h - yy - 3, 35, 8, MUTED)
    gap = cw / len(vals)
    for i, (v, lab) in enumerate(zip(vals, labels)):
        mid = cx0 + gap * (i + .5)
        yv = page_h - (cy0 + ch * (max_v - v) / (max_v - min_v))
        bottom = zero if v >= 0 else yv
        height = max(3, abs(yv-zero))
        canvas.setFillColor(col(CORAL if v < 0 else TEAL)); canvas.rect(mid - 28, bottom, 56, height, fill=1, stroke=0)
        value_y = yv + 7 if v >= 0 else bottom - 15
        canvas.setFont(bold, 9); canvas.setFillColor(col(NAVY)); canvas.drawCentredString(mid, value_y, f"{v:.1f}")
        wrapped(lab, mid - 45, 415, 90, 8, INK)
    wrapped("VAN (miles de USD) • 25 años • descuento 10% • supuestos del modelo vigente", 105, 472, 760, 9, MUTED)
    footer()

    # Slide 5: Q interpretation
    start("80 L/s en estiaje no significa 80 L/s todo el año", "04 / lectura hidrológica")
    card(60, 135, 405, 160, "Proxy de caudal de diseño", "Se limita la turbina a 80 L/s manteniendo la serie hipotética. VAN combinado con Hb 100 m y CAPEX +30%: −USD 64.7k.", GOLD, size=13)
    card(495, 135, 405, 160, "Límite severo de estrés", "Se supone 80 L/s en el río todo el año. Con 30 L/s reservados quedan 50 L/s turbinables; VAN combinado: −USD 192.9k.", CORAL, size=13)
    box(62, 335, 838, 90, "E5EFF3")
    wrapped("Para modelar el estiaje real faltan su duración, una curva de duración validada y el caudal ecológico aplicable.", 84, 356, 795, 15, NAVY, True)
    wrapped("El escenario constante es un límite de estrés, no una predicción.", 82, 444, 790, 11, MUTED)
    footer()

    # Slide 6: Gate0 work packages
    start("Gate 0: señales tempranas; campañas ampliadas opcionales", "05 / procura")
    card(52, 130, 270, 270, "A • Topografía", "Revisión y reconocimiento\nControl de campo\nHasta 8 ha para cotizar\nPolígono real antes de movilizar\n\nPuntos, perfiles, secciones y datos crudos", TEAL, size=11)
    card(345, 130, 270, 270, "B • Hidrología", "Auditoría de series\nReconocimiento de sección\n2–4 aforos iniciales\n\nOpción: 12 meses y 15–20 aforos, más análisis de disponibilidad y crecidas", GOLD, size=11)
    card(638, 130, 270, 270, "C • Geotecnia", "Revisión y recorrido\nPrograma de exploración cotizado\n\nOpción: 14 calicatas y 3 perforaciones iniciales, ajustadas por el especialista", CORAL, size=11)
    box(65, 430, 830, 45, MINT)
    wrapped("Pedir precios por lote, etapa, unidad, impuestos, movilización, plazo y entregables.", 82, 444, 800, 12, NAVY, True)
    footer()

    # Slide 7: restrictions
    start("Las restricciones todavía no están despejadas", "06 / control")
    risks = [("Hidráulica", "Hb < 95 m: revaluar", TEAL), ("Agua", "<70 L/s: alerta y validar", GOLD), ("Predial", "Sin servidumbre: variante", CORAL),
             ("Geotecnia", "Movimiento activo: costear", TEAL), ("Conexión", "Negativa: evaluar alternativa", GOLD), ("Ambiental", "Clasificación pendiente", CORAL)]
    for i, (title, body, accent) in enumerate(risks):
        x = 60 + (i % 3) * 290; top = 135 + (i // 3) * 118
        card(x, top, 265, 92, title, body, accent, size=10)
    box(62, 390, 838, 50, "FFF0D1")
    wrapped("95 m y 70 L/s son disparadores de revisión, no cancelación automática.", 82, 406, 800, 13, NAVY, True)
    wrapped("Estado: pendiente de polígono, datos de campo y respuestas institucionales.", 68, 462, 830, 9, MUTED)
    footer()

    # Slide 8: investment gates
    start("Financiar por etapas, con ofertas antes de avanzar", "07 / capital de desarrollo")
    stages = [("GATE 0", "USD 10–25k", "screening + reconocimiento", TEAL), ("GATE 1", "USD 30–60k", "factibilidad de sitio", GOLD), ("GATE 2", "USD 60–110k", "factibilidad cerrada / detalle", CORAL)]
    for i, (name, amount, scope, accent) in enumerate(stages):
        x = 60 + i * 292
        box(x, 155, 265, 175, WHITE)
        canvas.setFillColor(col(accent)); canvas.roundRect(x, page_h - 167, 265, 10, 5, fill=1, stroke=0)
        wrapped(name, x + 15, 188, 235, 12, accent, True)
        wrapped(amount, x + 12, 226, 240, 22, NAVY, True)
        wrapped(scope, x + 15, 270, 235, 10, MUTED)
    box(62, 372, 838, 78, "E5EFF3")
    wrapped("Rangos ROM de planificación: no son ofertas ni deben sumarse sin revisar solapes.", 83, 392, 795, 14, NAVY, True)
    wrapped("El ciclo hidrológico anual puede requerir fondos adicionales a Gate 0.", 82, 423, 790, 10, MUTED)
    footer()

    # Slide 9: approval
    start("Aprobar Gate 0 con entregables y criterios de salida", "08 / solicitud ejecutiva", True)
    wrapped("AUTORIZAR", 65, 132, 300, 11, CYAN, True)
    wrapped("Hasta USD 25,000", 65, 177, 480, 29, WHITE, True)
    wrapped("para obtener ofertas comparables, fijar el polígono, ejecutar reconocimiento y resolver incertidumbres críticas.", 66, 235, 435, 15, "D6E3EA", leading=21)
    card(525, 140, 375, 92, "1  Cotizar", "Tres lotes, base y opciones por separado.", TEAL, True, 11)
    card(525, 250, 375, 92, "2  Verificar", "Geometría, agua, acceso, predios y conexión.", GOLD, True, 11)
    card(525, 360, 375, 92, "3  Decidir", "Pasar a Gate 1 solo con ruta plausible y costos actualizados.", CORAL, True, 11)
    footer(True)

    # Slide 10: next actions
    start("Siguiente paso: cotizar y confirmar la ruta", "09 / acción inmediata")
    steps = ["Enviar el TDR a postores de topografía, hidrología y geotecnia.",
             "Pedir Gate 0 y opciones ampliadas por separado.",
             "Aprobar eje/polígono, acceso y responsable de datos crudos.",
             "Comparar ofertas y ejecutar el primer go/no-go."]
    for i, item in enumerate(steps):
        yy = 135 + i * 76
        wrapped(str(i+1), 68, yy, 36, 24, TEAL, True)
        wrapped(item, 130, yy + 4, 735, 15, NAVY, True)
    box(62, 430, 838, 48, NAVY)
    wrapped("Sin levantamiento real • sin aforo de campo • sin overlay GIS confirmado", 80, 445, 800, 12, WHITE, True)
    footer()
    canvas.save()


def main():
    pptx_path = ROOT / "Presentacion_Directiva_MCH_Rio_Blanco.pptx"
    add_presentation(pptx_path)
    add_pdf_presentation(ROOT / "Presentacion_Directiva_MCH_Rio_Blanco.pdf")
    md_to_docx(ROOT / "TDR_Estudios_Gate_0.md", ROOT / "TDR_Estudios_Gate_0.docx", "TDR de cotización Gate 0")
    md_to_docx(ROOT / "Decision_prefactibilidad_prueba_estres_TDR.md", ROOT / "Decision_Gate_0.docx", "Decisión Gate 0")
    profile = "file:///tmp/lo_mch_gate0_profile"
    subprocess.run([
        "libreoffice", "--headless", f"-env:UserInstallation={profile}",
        "--convert-to", "pdf", "--outdir", str(ROOT),
        str(ROOT / "TDR_Estudios_Gate_0.docx"), str(ROOT / "Decision_Gate_0.docx")
    ], check=True, capture_output=True, text=True)
    print("Generados:")
    for name in ["Presentacion_Directiva_MCH_Rio_Blanco.pptx", "Presentacion_Directiva_MCH_Rio_Blanco.pdf",
                 "TDR_Estudios_Gate_0.docx", "TDR_Estudios_Gate_0.pdf", "Decision_Gate_0.docx", "Decision_Gate_0.pdf"]:
        print(ROOT / name)


if __name__ == "__main__":
    main()
